from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from threegpp.models import (
    FieldProvenance,
    MetadataLayer,
    MetadataSourceKind,
    SourceArtifact,
    SpreadsheetParseResult,
    SpreadsheetParseSummary,
    TDocMetadata,
)

from .organizations import normalize_organizations


_COLUMN_ALIASES: dict[str, set[str]] = {
    "tdoc_id": {"tdoc", "tdoc id", "tdoc number", "tdoc no", "document id", "document number"},
    "title": {"title", "document title", "tdoc title"},
    "source_organization_raw": {
        "source", "source organization", "source organisation", "submitted by"
    },
    "agenda_item": {"agenda item", "agenda item number", "agenda", "ai"},
    "status": {"tdoc status", "status", "decision", "outcome"},
    "document_type": {"type", "document type", "tdoc type"},
    "document_category": {"category", "document category", "cr category"},
    "revision": {"revision", "revision number", "cr revision"},
    "abstract": {"abstract", "summary", "document abstract"},
    "agenda_item_description": {
        "agenda item description", "agenda description", "ai description"
    },
    "intended_for": {"for", "intended for", "purpose", "intended action"},
    "release": {"release", "3gpp release", "rel"},
    "specification": {"spec", "specification", "specification number", "ts tr"},
    "related_work_item": {
        "related wi", "related wis", "related work item", "related work items", "work item"
    },
}

_RELATED_ALIASES = {
    "is revision of",
    "revised to",
    "replaces",
    "replaced by",
    "related document",
    "related tdoc",
    "reply to",
    "original ls",
    "reply in",
}

_TDOC_ID = re.compile(r"\b([A-Z][A-Z0-9]?-\d{5,})\b", re.I)


class TDocListParser:
    def parse(self, path: Path, artifact: SourceArtifact) -> SpreadsheetParseResult:
        if artifact.retrieved_at is None or artifact.checksum is None:
            raise ValueError("TDoc-list parsing requires a downloaded artifact with checksum")
        sheets = self._read_sheets(path)
        sheet_name, rows, header_index, mapping, related_columns = self._select_layout(sheets)
        headers = [_cell_text(value) or f"column_{index + 1}" for index, value in enumerate(rows[header_index])]
        unknown_columns = [
            header
            for index, header in enumerate(headers)
            if index not in mapping.values() and index not in related_columns
        ]
        mapped_columns: dict[str, list[str]] = {
            canonical: [headers[index]] for canonical, index in mapping.items()
        }
        if related_columns:
            mapped_columns["related_tdoc_ids"] = [headers[index] for index in related_columns]

        tdocs: dict[str, TDocMetadata] = {}
        total_rows = unmapped = duplicates = 0
        for row in rows[header_index + 1 :]:
            if not any(_cell_text(value) for value in row):
                continue
            total_rows += 1
            raw_id = _value_at(row, mapping["tdoc_id"])
            identifier = _extract_primary_tdoc_id(raw_id)
            if identifier is None:
                unmapped += 1
                continue
            raw = {
                headers[index]: text
                for index, value in enumerate(row[: len(headers)])
                if (text := _cell_text(value)) is not None
            }
            values: dict[str, Any] = {}
            provenance: dict[str, FieldProvenance] = {}
            for field, index in mapping.items():
                if field == "tdoc_id":
                    continue
                raw_value = _value_at(row, index)
                value = (
                    _source_string(raw_value)
                    if field == "source_organization_raw"
                    else _cell_text(raw_value)
                )
                if value is None:
                    continue
                values[field] = value
                provenance[field] = FieldProvenance(
                    layer=MetadataLayer.TDOC_LIST,
                    source_url=artifact.source_url,
                    source_checksum=artifact.checksum,
                    source_column=headers[index],
                )
            related = _extract_related_ids(_value_at(row, index) for index in related_columns)
            if related:
                values["related_tdoc_ids"] = related
                provenance["related_tdoc_ids"] = FieldProvenance(
                    layer=MetadataLayer.TDOC_LIST,
                    source_url=artifact.source_url,
                    source_checksum=artifact.checksum,
                    source_column=", ".join(headers[index] for index in related_columns),
                )
            source_org = values.get("source_organization_raw")
            if source_org:
                values["organizations"] = normalize_organizations(source_org)
            record = TDocMetadata(
                tdoc_id=identifier,
                working_group=artifact.working_group,
                meeting=artifact.meeting,
                metadata_source_kind=MetadataSourceKind.TDOC_LIST,
                metadata_source_url=artifact.source_url,
                metadata_source_checksum=artifact.checksum,
                official_list_present=True,
                field_provenance=provenance,
                raw_metadata=raw,
                **values,
            )
            if record.tdoc_id in tdocs:
                duplicates += 1
            tdocs[record.tdoc_id] = record

        warnings: list[str] = []
        if unmapped:
            warnings.append(f"{unmapped} non-empty rows did not contain a recognized TDoc identifier")
        summary = SpreadsheetParseSummary(
            artifact_url=artifact.source_url,
            sheet_name=sheet_name,
            header_row=header_index + 1,
            total_data_rows=total_rows,
            mapped_rows=len(tdocs),
            unmapped_rows=unmapped,
            duplicate_rows=duplicates,
            mapped_columns=mapped_columns,
            unknown_columns=unknown_columns,
            warnings=warnings,
        )
        return SpreadsheetParseResult(summary=summary, tdocs=list(tdocs.values()))

    def _read_sheets(self, path: Path) -> dict[str, list[list[Any]]]:
        suffix = path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return {sheet.title: _bounded_xlsx_rows(sheet) for sheet in workbook.worksheets}
            finally:
                workbook.close()
        if suffix == ".xls":
            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                return {
                    sheet.name: [sheet.row_values(index) for index in range(sheet.nrows)]
                    for sheet in workbook.sheets()
                }
            finally:
                workbook.release_resources()
        raise ValueError(f"unsupported TDoc-list spreadsheet format: {path.suffix}")

    def _select_layout(
        self, sheets: dict[str, list[list[Any]]]
    ) -> tuple[str, list[list[Any]], int, dict[str, int], list[int]]:
        candidates: list[tuple[int, str, list[list[Any]], int, dict[str, int], list[int]]] = []
        for sheet_name, rows in sheets.items():
            for row_index, row in enumerate(rows[:25]):
                mapping, related = _map_headers(row)
                if "tdoc_id" in mapping:
                    candidates.append((len(mapping) + bool(related), sheet_name, rows, row_index, mapping, related))
        if not candidates:
            raise ValueError("no worksheet row contains a recognized TDoc identifier column")
        _, sheet_name, rows, row_index, mapping, related = max(
            candidates, key=lambda item: (item[0], -item[3], item[1].casefold())
        )
        return sheet_name, rows, row_index, mapping, related


def _bounded_xlsx_rows(sheet: Any, *, empty_limit: int = 100) -> list[list[Any]]:
    rows: list[list[Any]] = []
    empty_run = 0
    seen_data = False
    for values in sheet.iter_rows(values_only=True):
        row = list(values)
        has_data = any(_cell_text(value) is not None for value in row)
        if has_data:
            seen_data = True
            empty_run = 0
        elif seen_data:
            empty_run += 1
            if empty_run >= empty_limit:
                break
        rows.append(row)
    return rows


def _map_headers(row: list[Any]) -> tuple[dict[str, int], list[int]]:
    mapping: dict[str, int] = {}
    related: list[int] = []
    for index, value in enumerate(row):
        normalized = _normalize_header(_cell_text(value) or "")
        if normalized in _RELATED_ALIASES:
            related.append(index)
            continue
        for canonical, aliases in _COLUMN_ALIASES.items():
            if normalized in aliases and canonical not in mapping:
                mapping[canonical] = index
                break
    return mapping, related


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _source_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text.strip() else None


def _value_at(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def _extract_primary_tdoc_id(value: Any) -> str | None:
    text = _cell_text(value)
    if not text:
        return None
    match = _TDOC_ID.search(text.replace("_", "-"))
    return match.group(1).upper() if match else None


def _extract_related_ids(values: Any) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _cell_text(value)
        if not text:
            continue
        for match in _TDOC_ID.finditer(text.replace("_", "-")):
            identifier = match.group(1).upper()
            if identifier not in seen:
                seen.add(identifier)
                result.append(identifier)
    return result
