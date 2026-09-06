from __future__ import annotations

import io
import mimetypes
import zipfile
from abc import ABC, abstractmethod
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from .models import ContentBlock, ExtractionStatus

PARSER_VERSION = "1"
SUPPORTED_SUFFIXES = {".docx", ".pdf", ".xlsx", ".xls", ".txt"}


class DocumentParser(ABC):
    name = "base"
    version = PARSER_VERSION

    @abstractmethod
    def supports(self, filename: str) -> bool: ...

    @abstractmethod
    def parse(self, data: bytes, filename: str) -> tuple[list[ContentBlock], ExtractionStatus, list[str]]: ...

    def _ids(self, blocks: list[ContentBlock]) -> list[ContentBlock]:
        return [block.model_copy(update={"block_id": f"b{i:06d}", "order": i}) for i, block in enumerate(blocks, 1)]


class TextParser(DocumentParser):
    name = "text"
    def supports(self, filename: str) -> bool: return filename.lower().endswith(".txt")
    def parse(self, data: bytes, filename: str):
        text = data.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
        blocks = [ContentBlock(block_id="", type="paragraph", text=p.strip(), member_filename=filename, order=0) for p in text.split("\n\n") if p.strip()]
        status = ExtractionStatus.PARSED if blocks else ExtractionStatus.TEXT_UNAVAILABLE
        return self._ids(blocks), status, ([] if blocks else ["plain text contains no meaningful text"])


class DocxParser(DocumentParser):
    name = "docx"
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    def supports(self, filename: str) -> bool: return filename.lower().endswith(".docx")
    def parse(self, data: bytes, filename: str):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        blocks, headings = [], []
        body = root.find(f"{self.W}body")
        for node in list(body) if body is not None else []:
            if node.tag == f"{self.W}p":
                text = "".join(t.text or "" for t in node.iter(f"{self.W}t")).strip()
                if not text: continue
                style = node.find(f".//{self.W}pStyle")
                value = style.get(f"{self.W}val", "") if style is not None else ""
                level = int(value[-1]) if value.lower().startswith("heading") and value[-1:].isdigit() else None
                kind = "heading" if level else "paragraph"
                if level:
                    headings = headings[:level-1] + [text]
                blocks.append(ContentBlock(block_id="", type=kind, text=text, member_filename=filename, order=0, heading_level=level, heading_path=list(headings)))
            elif node.tag == f"{self.W}tbl":
                rows = [["".join(t.text or "" for t in cell.iter(f"{self.W}t")).strip() for cell in row.findall(f"{self.W}tc")] for row in node.findall(f"{self.W}tr")]
                blocks.append(ContentBlock(block_id="", type="table", member_filename=filename, order=0, rows=rows, heading_path=list(headings)))
        status = ExtractionStatus.PARSED if blocks else ExtractionStatus.TEXT_UNAVAILABLE
        return self._ids(blocks), status, ([] if blocks else ["DOCX contains no extractable text or tables"])


class PdfParser(DocumentParser):
    name = "pypdf"
    def supports(self, filename: str) -> bool: return filename.lower().endswith(".pdf")
    def parse(self, data: bytes, filename: str):
        from pypdf import PdfReader
        blocks = []
        for page_no, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(ContentBlock(block_id="", type="page", text=text, member_filename=filename, order=0, page_number=page_no))
        status = ExtractionStatus.PARSED if blocks else ExtractionStatus.TEXT_UNAVAILABLE
        return self._ids(blocks), status, ([] if blocks else ["PDF has no extractable text; OCR is not enabled"])


class SpreadsheetDocumentParser(DocumentParser):
    name = "spreadsheet"
    def supports(self, filename: str) -> bool: return filename.lower().endswith((".xlsx", ".xls"))
    def parse(self, data: bytes, filename: str):
        blocks = []
        if filename.lower().endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = ((ws.title, ws.iter_rows(values_only=True)) for ws in workbook.worksheets)
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            sheets = ((sheet.name, (sheet.row_values(i) for i in range(sheet.nrows))) for sheet in book.sheets())
        for name, rows in sheets:
            values = [[_cell(value) for value in row] for row in rows]
            while values and not any(value not in (None, "") for value in values[-1]): values.pop()
            blocks.append(ContentBlock(block_id="", type="sheet", member_filename=filename, order=0, sheet_name=name, rows=values))
        return self._ids(blocks), ExtractionStatus.PARSED, []


def _cell(value):
    if value is None or isinstance(value, (str, int, float, bool)): return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


PARSERS = [DocxParser(), PdfParser(), SpreadsheetDocumentParser(), TextParser()]


def parser_for(filename: str) -> DocumentParser | None:
    return next((parser for parser in PARSERS if parser.supports(filename)), None)


def media_type(filename: str) -> str:
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"
